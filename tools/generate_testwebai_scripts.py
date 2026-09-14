"""Complete the customer's TestWebAI command workbook for all three cases."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "TestWebAIScript1.xlsx"

CASES = [
    (
        "script1-student",
        [
            ("s", None, "1", "畫面等待1秒"),
            ("lf", "//html/body/div/form/div[1]/input", "apple", "First Name (姓名輸入)"),
            ("lf", "//html/body/div/form/div[2]/input", "test@example.com", "E-Mail (電子郵件)"),
            ("lseloptlc", "//html/body/div/form/div[3]/select", "student", "身分角色(學生)"),
            ("lf", "//html/body/div/form/div[4]/input", "台灣大學", "學校名稱"),
            ("lc", "//html/body/div/form/div[5]/div/label[2]/input", "1", "性別 (女)"),
            ("lc", "//html/body/div/form/div[6]/div/label[3]/input", "1", "興趣 (看電影)"),
            ("lc", "//html/body/div/form/div[7]/input", "1", "提交表單"),
            ("E", None, None, None),
        ],
    ),
    (
        "script2-developer",
        [
            ("s", None, "1", "畫面等待1秒"),
            ("lf", "//html/body/div/form/div[1]/input", "charles", "First Name (姓名輸入)"),
            ("lf", "//html/body/div/form/div[2]/input", "dev.charles@example.com", "E-Mail (電子郵件)"),
            ("lseloptlc", "//html/body/div/form/div[3]/select", "developer", "身分角色(工程師)"),
            ("lc", "//html/body/div/form/div[5]/div/label[1]/input", "1", "性別 (男)"),
            ("lc", "//html/body/div/form/div[6]/div/label[1]/input", "1", "興趣 (寫程式)"),
            ("lc", "//html/body/div/form/div[7]/input", "1", "提交表單"),
            ("E", None, None, None),
        ],
    ),
    (
        "script3-designer",
        [
            ("s", None, "1", "畫面等待1秒"),
            ("lf", "//html/body/div/form/div[1]/input", "david", "First Name (姓名輸入)"),
            ("lf", "//html/body/div/form/div[2]/input", "designer.d@example.com", "E-Mail (電子郵件)"),
            ("lseloptlc", "//html/body/div/form/div[3]/select", "designer", "身分角色(設計師)"),
            ("lc", "//html/body/div/form/div[5]/div/label[1]/input", "1", "性別 (男)"),
            ("lc", "//html/body/div/form/div[6]/div/label[2]/input", "1", "興趣 (打電動)"),
            ("lc", "//html/body/div/form/div[6]/div/label[3]/input", "1", "興趣 (看電影)"),
            ("lc", "//html/body/div/form/div[7]/input", "1", "提交表單"),
            ("E", None, None, None),
        ],
    ),
]


def copy_sheet_layout(source, target) -> None:
    for row in source.iter_rows():
        for cell in row:
            new_cell = target[cell.coordinate]
            if cell.has_style:
                new_cell._style = copy(cell._style)
            if cell.number_format:
                new_cell.number_format = cell.number_format
            if cell.alignment:
                new_cell.alignment = copy(cell.alignment)
            if cell.protection:
                new_cell.protection = copy(cell.protection)
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
        target.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
        target.row_dimensions[key].hidden = dim.hidden


def write_case(ws, commands) -> None:
    for row in range(2, 11):
        for column in range(1, 14):
            ws.cell(row, column).value = None
    for row, (code, locator, value, note) in enumerate(commands, start=2):
        ws.cell(row, 1).value = code
        ws.cell(row, 2).value = locator
        ws.cell(row, 5).value = value
        ws.cell(row, 13).value = note


def main() -> None:
    workbook = load_workbook(OUTPUT)
    template = workbook[workbook.sheetnames[0]]
    for ws in workbook.worksheets[1:]:
        workbook.remove(ws)
    template.title = CASES[0][0]
    write_case(template, CASES[0][1])

    for title, commands in CASES[1:]:
        ws = workbook.create_sheet(title)
        copy_sheet_layout(template, ws)
        for row in template.iter_rows():
            for cell in row:
                ws[cell.coordinate].value = cell.value
        write_case(ws, commands)

    workbook.save(OUTPUT)


if __name__ == "__main__":
    main()
